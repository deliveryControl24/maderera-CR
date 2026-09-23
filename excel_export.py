"""Exportacion a Excel (.xlsx) para reportes e inventario."""
import os
from datetime import datetime


def exportar_excel(archivo, encabezados, filas, titulo=None, hoja="Reporte"):
    """
    Guarda datos en un archivo .xlsx.
    encabezados: lista de textos de columna
    filas: lista de listas/tuplas con los valores
    Retorna (ok: bool, mensaje: str)
    """
    if not archivo:
        return False, "Ruta de archivo vacia."
    if not str(archivo).lower().endswith(".xlsx"):
        archivo = str(archivo) + ".xlsx"

    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        from openpyxl.utils import get_column_letter
    except ImportError:
        return _exportar_csv_fallback(archivo, encabezados, filas)

    try:
        wb = Workbook()
        ws = wb.active
        ws.title = (hoja or "Reporte")[:31]

        start_row = 1
        if titulo:
            ws.cell(row=1, column=1, value=titulo)
            ws.cell(row=1, column=1).font = Font(bold=True, size=13, color="1B5E20")
            ws.merge_cells(start_row=1, start_column=1, end_row=1,
                           end_column=max(1, len(encabezados)))
            ws.cell(row=2, column=1, value=f"Generado: {datetime.now().strftime('%Y-%m-%d %H:%M')}")
            ws.cell(row=2, column=1).font = Font(size=9, italic=True, color="666666")
            start_row = 4

        header_fill = PatternFill("solid", fgColor="1B5E20")
        header_font = Font(bold=True, color="FFFFFF", size=10)
        thin = Border(
            left=Side(style="thin", color="B0BEC5"),
            right=Side(style="thin", color="B0BEC5"),
            top=Side(style="thin", color="B0BEC5"),
            bottom=Side(style="thin", color="B0BEC5"),
        )
        alt_fill = PatternFill("solid", fgColor="E8F5E9")

        for col, h in enumerate(encabezados, 1):
            c = ws.cell(row=start_row, column=col, value=str(h))
            c.fill = header_fill
            c.font = header_font
            c.alignment = Alignment(horizontal="center", wrap_text=True)
            c.border = thin

        for i, fila in enumerate(filas):
            r = start_row + 1 + i
            for col, val in enumerate(fila, 1):
                c = ws.cell(row=r, column=col, value=val)
                c.border = thin
                if i % 2 == 1:
                    c.fill = alt_fill
                if isinstance(val, (int, float)) and not isinstance(val, bool):
                    c.number_format = "#,##0.00"

        # Ancho de columnas (aprox)
        for col in range(1, len(encabezados) + 1):
            max_len = len(str(encabezados[col - 1] or ""))
            for fila in filas[:80]:
                if col - 1 < len(fila):
                    max_len = max(max_len, len(str(fila[col - 1] if fila[col - 1] is not None else "")))
            ws.column_dimensions[get_column_letter(col)].width = min(40, max(10, max_len + 2))

        ws.auto_filter.ref = (
            f"A{start_row}:{get_column_letter(len(encabezados))}"
            f"{start_row + len(filas)}"
        )
        ws.freeze_panes = ws.cell(row=start_row + 1, column=1)

        os.makedirs(os.path.dirname(os.path.abspath(archivo)) or ".", exist_ok=True)
        wb.save(archivo)
        return True, archivo

    except Exception as e:
        return False, str(e)


def _exportar_csv_fallback(archivo, encabezados, filas):
    """Si openpyxl no esta disponible, guarda CSV con nombre .csv."""
    try:
        csv_path = os.path.splitext(archivo)[0] + ".csv"
        import csv
        with open(csv_path, "w", newline="", encoding="utf-8-sig") as f:
            w = csv.writer(f)
            w.writerow(encabezados)
            for fila in filas:
                w.writerow(fila)
        return True, csv_path
    except Exception as e:
        return False, str(e)


def tree_a_excel(tree, archivo, titulo=None, hoja="Datos"):
    """Exporta un Treeview completo a Excel."""
    if not tree.get_children():
        return False, "No hay datos para exportar."
    columnas = list(tree["columns"])
    encabezados = [tree.heading(c)["text"] for c in columnas]
    filas = []
    for item in tree.get_children():
        vals = tree.item(item)["values"]
        filas.append(list(vals) if vals else [])
    return exportar_excel(archivo, encabezados, filas, titulo=titulo, hoja=hoja)
